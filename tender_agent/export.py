from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from tender_agent.outline import Outline
from tender_agent.textutil import safe_filename


def _set_run_font(run, size: int, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")


def export_markdown(markdown: str, directory: Path, stem: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{stem}.md"
    path.write_text(markdown, encoding="utf-8")
    return path


def export_docx(outline: Outline, directory: Path, stem: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(outline.title)
    _set_run_font(run, 22, bold=True)

    toc_heading = document.add_heading("目录", level=1)
    for paragraph in (toc_heading,):
        for text_run in paragraph.runs:
            _set_run_font(text_run, 16, bold=True)
    for section in outline.sections:
        item = document.add_paragraph()
        item.paragraph_format.left_indent = Pt(12 * max(section.level - 1, 0))
        text_run = item.add_run(f"{section.number} {section.title}")
        _set_run_font(text_run, 12)

    for section in outline.sections:
        heading = document.add_heading(f"{section.number} {section.title}", level=min(section.level, 3))
        for text_run in heading.runs:
            _set_run_font(text_run, 16 if section.level == 1 else 14, bold=True)
        for block in _blocks(section.content or ""):
            if block.startswith("###"):
                sub = document.add_heading(re.sub(r"^#+\s*", "", block), level=min(section.level + 1, 3))
                for text_run in sub.runs:
                    _set_run_font(text_run, 13, bold=True)
                continue
            paragraph = document.add_paragraph()
            if block.startswith("- "):
                text = block[2:].strip()
                try:
                    paragraph.style = document.styles["List Bullet"]
                except KeyError:
                    text = f"• {text}"
            else:
                text = block
            text_run = paragraph.add_run(text)
            _set_run_font(text_run, 12)

    path = directory / f"{stem}.docx"
    document.save(path)
    return path


def _blocks(content: str) -> list[str]:
    blocks: list[str] = []
    for raw in re.split(r"\n+", content):
        line = raw.strip()
        if not line or line == "---":
            continue
        blocks.append(line)
    return blocks


def make_stem(project_name: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"{safe_filename(project_name)}-{stamp}"
